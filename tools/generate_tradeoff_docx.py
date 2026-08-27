from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Inches, Pt


OUTPUT_PATH = "docs/AFN_TRADE_OFF_ANALYSIS_SAMPLE_STYLE_REVISED.docx"


ROWS = [
    ("1. User Registration", "Fast account creation", "High", "Low cost"),
    ("2. Login / Authentication", "Fast login and verification", "High", "Low cost"),
    ("3. Password Reset", "Quick reset process", "Medium-High", "Low cost"),
    ("4. Client Service Request", "Fast request submission", "High", "Low cost"),
    ("5. Service Approval", "Fast admin approval", "High", "Low cost"),
    ("6. Ticket Creation", "Fast ticket generation", "High", "Low cost"),
    ("7. Technician Assignment", "Fast job assignment", "Medium-High", "Medium cost"),
    ("8. Auto Dispatch", "Fast technician matching", "Medium-High", "Medium cost"),
    ("9. Technician Dashboard", "Fast access to assigned jobs", "High", "Low cost"),
    ("10. Navigation / Arrival", "Fast location validation", "Medium-High", "Medium cost"),
    ("11. Inspection Checklist", "Fast checklist encoding", "High", "Low cost"),
    ("12. Work Status Updates", "Fast progress tracking", "High", "Low cost"),
    ("13. Inventory Management", "Fast basic stock monitoring", "High", "Low cost"),
    ("14. Inventory Customization", "Limited item specification editing", "Medium", "Medium cost"),
    ("15. Parts Request", "Fast parts reservation", "High", "Medium cost"),
    ("16. Document Generation", "Fast report creation", "Medium-High", "Low cost"),
    ("17. Warranty Tracking", "Fast warranty checking", "High", "Low cost"),
    ("18. After-Sales Support", "Fast support monitoring", "Medium-High", "Low cost"),
    ("19. Messaging / Communication", "Fast sending and receiving", "Medium-High", "Low cost"),
    ("20. Reports and Analytics", "Fast system monitoring", "High", "Low cost"),
    ("21. API / Database Services", "Fast storage and retrieval", "High", "Low cost"),
]


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table):
    table.style = "Table Grid"
    widths = [Inches(2.7), Inches(2.8), Inches(1.35), Inches(1.45)]
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Trade-off Analysis")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)

    table = doc.add_table(rows=1, cols=4)
    headers = ["System Component", "Performance", "Security", "Affordability"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True)

    for row_data in ROWS:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            set_cell_text(row.cells[idx], value)

    style_table(table)

    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    main()
